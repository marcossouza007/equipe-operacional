"""
Módulo de Upload - Processa múltiplos formatos de arquivo
"""
import os
import asyncio
from typing import List, Dict, Any, Optional
from pathlib import Path
import aiofiles
from fastapi import UploadFile, File
import PyPDF2
from docx import Document
import openpyxl
import csv
import json
import yaml
from .file_classifier import FileClassifier, FileCategory


class FileUploader:
    """Gerenciador de upload de arquivos"""
    
    ALLOWED_EXTENSIONS = {
        '.pdf', '.docx', '.xlsx', '.csv', '.txt', '.json', '.xml', 
        '.yml', '.yaml', '.log', '.conf', '.config', '.png', '.jpg', '.jpeg'
    }
    
    def __init__(self, upload_dir: str = "uploads"):
        self.upload_dir = Path(upload_dir)
        self.upload_dir.mkdir(exist_ok=True)
        self.classifier = FileClassifier()
        self.processed_files: Dict[str, Dict[str, Any]] = {}
    
    async def upload_file(self, file: UploadFile) -> Dict[str, Any]:
        """
        Realiza upload e processa arquivo
        
        Args:
            file: Arquivo enviado
            
        Returns:
            Informações do arquivo processado
        """
        # Validar extensão
        file_ext = Path(file.filename).suffix.lower()
        if file_ext not in self.ALLOWED_EXTENSIONS:
            return {
                "success": False,
                "error": f"Tipo de arquivo não permitido: {file_ext}",
                "allowed_types": list(self.ALLOWED_EXTENSIONS)
            }
        
        # Salvar arquivo
        file_path = self.upload_dir / file.filename
        async with aiofiles.open(file_path, 'wb') as f:
            content = await file.read()
            await f.write(content)
        
        # Processar arquivo
        try:
            file_content = await self._extract_content(file_path, file_ext)
            category, confidence = self.classifier.classify_file(file.filename, file_content[:1000])
            
            file_info = {
                "success": True,
                "filename": file.filename,
                "file_path": str(file_path),
                "file_size": len(content),
                "file_type": file_ext,
                "category": category.value,
                "confidence": confidence,
                "content_preview": file_content[:500],
                "content_length": len(file_content)
            }
            
            self.processed_files[file.filename] = file_info
            return file_info
            
        except Exception as e:
            return {
                "success": False,
                "filename": file.filename,
                "error": str(e)
            }
    
    async def upload_multiple(self, files: List[UploadFile]) -> List[Dict[str, Any]]:
        """
        Realiza upload de múltiplos arquivos
        
        Args:
            files: Lista de arquivos
            
        Returns:
            Lista de informações dos arquivos processados
        """
        results = []
        for file in files:
            result = await self.upload_file(file)
            results.append(result)
        return results
    
    async def _extract_content(self, file_path: Path, file_ext: str) -> str:
        """
        Extrai conteúdo do arquivo
        
        Args:
            file_path: Caminho do arquivo
            file_ext: Extensão do arquivo
            
        Returns:
            Conteúdo extraído
        """
        if file_ext == '.pdf':
            return await self._extract_pdf(file_path)
        elif file_ext == '.docx':
            return await self._extract_docx(file_path)
        elif file_ext == '.xlsx':
            return await self._extract_xlsx(file_path)
        elif file_ext == '.csv':
            return await self._extract_csv(file_path)
        elif file_ext == '.json':
            return await self._extract_json(file_path)
        elif file_ext == '.xml':
            return await self._extract_xml(file_path)
        elif file_ext in ['.yml', '.yaml']:
            return await self._extract_yaml(file_path)
        else:
            return await self._extract_text(file_path)
    
    async def _extract_pdf(self, file_path: Path) -> str:
        """Extrai texto de PDF"""
        text = ""
        try:
            with open(file_path, 'rb') as f:
                pdf_reader = PyPDF2.PdfReader(f)
                for page in pdf_reader.pages:
                    text += page.extract_text() + "\n"
        except Exception as e:
            text = f"Erro ao ler PDF: {str(e)}"
        return text
    
    async def _extract_docx(self, file_path: Path) -> str:
        """Extrai texto de DOCX"""
        text = ""
        try:
            doc = Document(file_path)
            for para in doc.paragraphs:
                text += para.text + "\n"
        except Exception as e:
            text = f"Erro ao ler DOCX: {str(e)}"
        return text
    
    async def _extract_xlsx(self, file_path: Path) -> str:
        """Extrai conteúdo de XLSX"""
        text = ""
        try:
            wb = openpyxl.load_workbook(file_path)
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                text += f"\n=== Sheet: {sheet_name} ===\n"
                for row in ws.iter_rows(values_only=True):
                    text += str(row) + "\n"
        except Exception as e:
            text = f"Erro ao ler XLSX: {str(e)}"
        return text
    
    async def _extract_csv(self, file_path: Path) -> str:
        """Extrai conteúdo de CSV"""
        text = ""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                reader = csv.reader(f)
                for row in reader:
                    text += str(row) + "\n"
        except Exception as e:
            text = f"Erro ao ler CSV: {str(e)}"
        return text
    
    async def _extract_json(self, file_path: Path) -> str:
        """Extrai conteúdo de JSON"""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                data = json.load(f)
                return json.dumps(data, indent=2, ensure_ascii=False)
        except Exception as e:
            return f"Erro ao ler JSON: {str(e)}"
    
    async def _extract_xml(self, file_path: Path) -> str:
        """Extrai conteúdo de XML"""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                return f.read()
        except Exception as e:
            return f"Erro ao ler XML: {str(e)}"
    
    async def _extract_yaml(self, file_path: Path) -> str:
        """Extrai conteúdo de YAML"""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                data = yaml.safe_load(f)
                return json.dumps(data, indent=2, ensure_ascii=False)
        except Exception as e:
            return f"Erro ao ler YAML: {str(e)}"
    
    async def _extract_text(self, file_path: Path) -> str:
        """Extrai conteúdo de arquivo de texto"""
        try:
            async with aiofiles.open(file_path, 'r', encoding='utf-8') as f:
                return await f.read()
        except Exception as e:
            return f"Erro ao ler arquivo: {str(e)}"
    
    def get_processed_files(self) -> Dict[str, Dict[str, Any]]:
        """Retorna lista de arquivos processados"""
        return self.processed_files
    
    def get_files_by_category(self, category: FileCategory) -> List[Dict[str, Any]]:
        """Retorna arquivos filtrados por categoria"""
        return [f for f in self.processed_files.values() if f.get('category') == category.value]
